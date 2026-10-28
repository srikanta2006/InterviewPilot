import io
import pytest
from app.services.parser import (
    CorruptDocumentError,
    DocumentParsingError,
    EmptyDocumentError,
    FileTooLargeError,
    ScannedDocumentError,
    UnsupportedDocumentError,
    parse_document,
    parse_docx,
    parse_pdf,
    parse_txt,
)


def test_parse_txt_valid():
    sample_text = "John Doe is a Senior Software Engineer with 7 years of Python, FastAPI, and React expertise."
    result = parse_txt(sample_text.encode("utf-8"))
    assert "John Doe" in result
    assert "FastAPI" in result


def test_parse_txt_empty():
    with pytest.raises(EmptyDocumentError):
        parse_txt(b"")

    with pytest.raises(EmptyDocumentError):
        parse_txt(b"Too short")


def test_parse_document_unsupported_ext():
    with pytest.raises(UnsupportedDocumentError) as exc_info:
        parse_document(b"sample bytes" * 10, "resume.exe")
    assert "Unsupported file format" in str(exc_info.value)


def test_parse_document_legacy_doc():
    with pytest.raises(UnsupportedDocumentError) as exc_info:
        parse_document(b"sample bytes" * 10, "resume.doc")
    assert "Legacy .doc format is not supported" in str(exc_info.value)


def test_parse_document_too_large():
    huge_bytes = b"0" * (11 * 1024 * 1024)
    with pytest.raises(FileTooLargeError):
        parse_document(huge_bytes, "resume.pdf")


def test_parse_pdf_valid():
    import pymupdf as fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (50, 50),
        "Alex Rivera - Senior Full Stack Developer\nExperienced in Python, FastAPI, React, and Cloud Systems."
    )
    pdf_bytes = doc.tobytes()
    doc.close()

    result = parse_pdf(pdf_bytes)
    assert "Alex Rivera" in result
    assert "FastAPI" in result


def test_parse_pdf_scanned_or_empty():
    import pymupdf as fitz
    doc = fitz.open()
    doc.new_page()  # blank page without text
    pdf_bytes = doc.tobytes()
    doc.close()

    with pytest.raises(ScannedDocumentError):
        parse_pdf(pdf_bytes)


def test_parse_pdf_corrupt():
    with pytest.raises(CorruptDocumentError):
        parse_pdf(b"Malformed non-pdf byte sequence")


def test_parse_docx_valid():
    import docx
    doc = docx.Document()
    doc.add_paragraph("Sarah Connor - Systems Architect and Python Specialist.")
    doc.add_paragraph("Core competencies: Distributed Systems, FastAPI, Docker, and Kubernetes.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    result = parse_docx(docx_bytes)
    assert "Sarah Connor" in result
    assert "Distributed Systems" in result


def test_parse_docx_empty():
    import docx
    doc = docx.Document()
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    with pytest.raises(EmptyDocumentError):
        parse_docx(docx_bytes)


def test_parse_docx_corrupt():
    with pytest.raises(CorruptDocumentError):
        parse_docx(b"PK\x03\x04 corrupt zip payload")
