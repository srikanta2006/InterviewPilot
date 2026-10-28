import io
import os
from typing import Optional

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB
MIN_EXTRACTED_CHARS = 30  # Threshold to detect empty or scanned image documents


class DocumentParsingError(Exception):
    """Base exception for document parsing failures."""
    pass


class UnsupportedDocumentError(DocumentParsingError):
    """Raised when file format is not supported."""
    pass


class FileTooLargeError(DocumentParsingError):
    """Raised when file exceeds maximum allowed size."""
    pass


class EmptyDocumentError(DocumentParsingError):
    """Raised when document content is empty."""
    pass


class ScannedDocumentError(DocumentParsingError):
    """Raised when document appears to be scanned or image-only with insufficient text."""
    pass


class CorruptDocumentError(DocumentParsingError):
    """Raised when document structure cannot be read or parsed."""
    pass


def parse_pdf(file_bytes: bytes) -> str:
    """Extract readable text from a PDF file using PyMuPDF."""
    try:
        import pymupdf as fitz  # PyMuPDF
    except ImportError as e:
        raise DocumentParsingError("PyMuPDF is not installed in the environment.") from e

    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise CorruptDocumentError(f"Failed to read PDF document: {e}") from e

    if doc.page_count == 0:
        raise EmptyDocumentError("PDF document contains no pages.")

    extracted_pages = []
    for page_idx in range(doc.page_count):
        page = doc.load_page(page_idx)
        text = page.get_text()
        if text.strip():
            extracted_pages.append(text.strip())

    full_text = "\n\n".join(extracted_pages).strip()
    if not full_text or len(full_text) < MIN_EXTRACTED_CHARS:
        raise ScannedDocumentError(
            "PDF contains little or no extractable text. Scanned or image-only documents are not supported. "
            "Please provide a text-based PDF or paste text directly."
        )

    return full_text


def parse_docx(file_bytes: bytes) -> str:
    """Extract readable text and tables from a DOCX file using python-docx."""
    try:
        import docx
    except ImportError as e:
        raise DocumentParsingError("python-docx is not installed in the environment.") from e

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise CorruptDocumentError(f"Failed to read DOCX document: {e}") from e

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    table_lines = []
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                table_lines.append(row_text)

    all_text = "\n\n".join(paragraphs + table_lines).strip()
    if not all_text or len(all_text) < MIN_EXTRACTED_CHARS:
        raise EmptyDocumentError("DOCX document contains no readable text content.")

    return all_text


def parse_txt(file_bytes: bytes) -> str:
    """Extract text from a plain text file."""
    try:
        text = file_bytes.decode("utf-8").strip()
    except UnicodeDecodeError:
        try:
            text = file_bytes.decode("latin-1").strip()
        except Exception as e:
            raise CorruptDocumentError(f"Failed to decode text file: {e}") from e

    if not text or len(text) < MIN_EXTRACTED_CHARS:
        raise EmptyDocumentError("Text document contains insufficient or empty text.")

    return text


def parse_document(file_bytes: bytes, filename: str) -> str:
    """Validate file metadata and extract textual content."""
    if not file_bytes:
        raise EmptyDocumentError(f"Uploaded file '{filename}' is empty (0 bytes).")

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        raise FileTooLargeError(
            f"File '{filename}' exceeds maximum allowed size of {MAX_FILE_SIZE_BYTES // (1024 * 1024)}MB."
        )

    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return parse_pdf(file_bytes)
    elif ext == ".docx":
        return parse_docx(file_bytes)
    elif ext == ".doc":
        raise UnsupportedDocumentError("Legacy .doc format is not supported. Please convert to modern .docx or .pdf.")
    elif ext == ".txt":
        return parse_txt(file_bytes)
    else:
        raise UnsupportedDocumentError(
            f"Unsupported file format '{ext}'. Supported formats: .pdf, .docx, .txt"
        )
